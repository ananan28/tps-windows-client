"""Independent local-page extraction and input validation."""
import json
import re
from datetime import datetime, timezone
from urllib.parse import urlsplit

VERSION = '0.1.11'
HOME = 'https://www.truepeoplesearch.com/'


def proxy_config(raw):
    raw = raw.strip()
    if not raw:
        return None
    if '://' in raw:
        u = urlsplit(raw)
        if u.scheme != 'http' or not u.hostname or not u.port:
            raise ValueError('请输入 HTTP 代理，或 host:port:user:password')
        if u.path not in ('', '/') or u.query or u.fragment:
            raise ValueError('代理地址不能包含路径或查询参数')
        from urllib.parse import unquote
        return {'server': f'http://{u.hostname}:{u.port}',
                'username': unquote(u.username or ''), 'password': unquote(u.password or '')}
    parts = raw.split(':', 3)
    if len(parts) not in (2, 4):
        raise ValueError('代理格式应为 host:port 或 host:port:user:password')
    host, port = parts[:2]
    if not re.fullmatch(r'[A-Za-z0-9.-]+', host) or not port.isdigit() or not 1 <= int(port) <= 65535:
        raise ValueError('代理地址或端口不正确')
    result = {'server': f'http://{host}:{port}'}
    if len(parts) == 4:
        if not parts[2] or not parts[3]:
            raise ValueError('代理账号和密码不能为空')
        result.update(username=parts[2], password=parts[3])
    return result


def query_value(kind, raw):
    raw = raw.strip()
    if kind == '手机号':
        digits = re.sub(r'\D', '', raw)
        if len(digits) == 11 and digits.startswith('1'):
            digits = digits[1:]
        if len(digits) != 10:
            raise ValueError('请输入 10 位美国手机号，或以 1 开头的 11 位号码')
        return digits
    if kind != '邮箱' or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', raw):
        raise ValueError('请输入有效邮箱')
    return raw


def site_url(url):
    u = urlsplit(url)
    return u.scheme == 'https' and u.hostname in ('truepeoplesearch.com', 'www.truepeoplesearch.com') and u.port in (None, 443)


def blocked(text, url=''):
    text = text.lower()
    return any(x in text for x in ('verify you are human', 'captcha challenge',
        'sorry, you have been blocked', 'access denied', 'rate limited',
        'verification required', 'request failed, please go back')) or any(
        x in url.lower() for x in ('internalcaptcha', 'ratelimited', '/challenge'))


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def extract(snapshot):
    """Only extract a confirmed detail page, never a result listing/challenge."""
    url = snapshot.get('url', '')
    if not site_url(url):
        raise ValueError('只允许读取 TruePeopleSearch 官网详情页')
    if blocked(snapshot.get('title', '') + ' ' + snapshot.get('text', ''), url):
        raise ValueError('网站要求验证或限制访问；请在浏览器人工处理，结果未保存')
    if not snapshot.get('is_detail'):
        raise ValueError('当前不是可识别的详情页，请先点击 View Details')
    person = {}
    def visit(obj):
        nonlocal person
        if isinstance(obj, list):
            for v in obj: visit(v)
        elif isinstance(obj, dict):
            typ = obj.get('@type', [])
            if typ == 'Person' or isinstance(typ, list) and 'Person' in typ:
                if not person: person = obj
            if '@graph' in obj: visit(obj['@graph'])
    for item in snapshot.get('jsonld', []):
        try: visit(json.loads(item))
        except (ValueError, TypeError): pass
    def values(v):
        return v if isinstance(v, list) else [v] if v else []
    addresses = []
    for a in values(person.get('address')):
        if isinstance(a, dict):
            addresses.append(clean(', '.join(clean(a.get(k)) for k in
                ('streetAddress', 'addressLocality', 'addressRegion', 'postalCode') if a.get(k))))
        else: addresses.append(clean(a))
    phones = values(person.get('telephone')) + snapshot.get('phones', [])
    emails = values(person.get('email')) + snapshot.get('emails', [])
    name = clean(person.get('name') or snapshot.get('name'))
    if not name:
        raise ValueError('未识别到姓名；网页结构可能已变化，未保存空结果')
    unique = lambda xs: ' | '.join(dict.fromkeys(clean(x) for x in xs if clean(x)))
    return {'name': name, 'age': clean(snapshot.get('age')),
            'addresses': unique(addresses + snapshot.get('addresses', [])),
            'phones': unique(phones), 'emails': unique(emails), 'source_url': url,
            'collected_at': datetime.now(timezone.utc).isoformat(timespec='seconds')}


def export_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Results'
    keys = ['query', 'name', 'age', 'addresses', 'phones', 'emails', 'source_url', 'collected_at']
    ws.append(keys)
    for row in rows:
        ws.append([str(row.get(k, ''))[:32767] for k in keys])
        for cell in ws[ws.max_row]:
            cell.data_type = 's'  # Untrusted page text must never become a spreadsheet formula.
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(color='FFFFFF', bold=True)
        c.fill = PatternFill('solid', fgColor='2457A7')
    for col in ('A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'):
        ws.column_dimensions[col].width = 28 if col not in ('D', 'G') else 55
    import os, tempfile
    from pathlib import Path
    dest = Path(path)
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=dest.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, dest)
    finally:
        if os.path.exists(tmp): os.unlink(tmp)


